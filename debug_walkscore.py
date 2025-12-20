#!/usr/bin/env python3

import os
from openpyxl import load_workbook
import glob

def debug_walkscore_issue():
    """Debug why walkscore values are not displayed in Excel"""
    
    # Find Excel files
    excel_files = glob.glob('*.xlsx') + glob.glob('extraction-*.xlsx')
    
    if not excel_files:
        print("❌ No Excel files found!")
        print("   Run extract-centris-data.py first to create the Excel file")
        return
    
    for excel_file in excel_files:
        print(f"\n📊 Analyzing {excel_file}...")
        
        try:
            wb = load_workbook(excel_file)
            ws = wb.active
            
            # Find WalkScore column
            walkscore_col = None
            address_col = None
            
            for col in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=col).value
                if header == "WalkScore":
                    walkscore_col = col
                elif header == "Adresse":
                    address_col = col
            
            if not walkscore_col:
                print("❌ No WalkScore column found!")
                continue
                
            print(f"✅ WalkScore column found at position {walkscore_col}")
            
            # Check data
            total_rows = ws.max_row - 1  # Exclude header
            empty_walkscores = 0
            filled_walkscores = 0
            
            print(f"\n📋 Checking {total_rows} listings:")
            
            for row in range(2, min(12, ws.max_row + 1)):  # Check first 10 rows
                address = ws.cell(row=row, column=address_col).value if address_col else "N/A"
                walkscore = ws.cell(row=row, column=walkscore_col).value
                
                if walkscore is None or walkscore == "":
                    empty_walkscores += 1
                    status = "❌ EMPTY"
                else:
                    filled_walkscores += 1
                    status = f"✅ {walkscore}"
                
                print(f"   Row {row}: {status} | {address}")
            
            # Summary
            for row in range(12, ws.max_row + 1):
                walkscore = ws.cell(row=row, column=walkscore_col).value
                if walkscore is None or walkscore == "":
                    empty_walkscores += 1
                else:
                    filled_walkscores += 1
            
            print(f"\n📊 Summary:")
            print(f"   Total listings: {total_rows}")
            print(f"   Empty WalkScores: {empty_walkscores}")
            print(f"   Filled WalkScores: {filled_walkscores}")
            
            if empty_walkscores > 0:
                print(f"\n💡 Solution: Run extended-parse.py to fill missing WalkScore values")
                print(f"   This will fetch WalkScore data for {empty_walkscores} listings")
            
        except Exception as e:
            print(f"❌ Error reading {excel_file}: {e}")

if __name__ == "__main__":
    debug_walkscore_issue()
