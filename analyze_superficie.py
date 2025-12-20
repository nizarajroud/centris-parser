#!/usr/bin/env python3

import os
from openpyxl import load_workbook
from dotenv import load_dotenv

load_dotenv()

def analyze_superficie_issue():
    """Analyze SuperficieDuterrain values issue in Excel file"""
    
    excel_file = os.getenv('EXTRACTION_CENTRIS', 'extraction-centris.xlsx')
    
    if not os.path.exists(excel_file):
        print(f"❌ Excel file not found: {excel_file}")
        return
    
    print(f"📊 Analyzing SuperficieDuterrain in {excel_file}...")
    
    wb = load_workbook(excel_file)
    ws = wb.active
    
    # Find columns
    centris_col = None
    superficie_col = None
    address_col = None
    
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header == "No Centris":
            centris_col = col
        elif header == "SuperficieDuterrain":
            superficie_col = col
        elif header == "Adresse":
            address_col = col
    
    print(f"✅ Found columns - Centris: {centris_col}, SuperficieDuterrain: {superficie_col}, Address: {address_col}")
    
    # Analyze the data
    total_rows = ws.max_row - 1
    filled_superficie = 0
    empty_superficie = 0
    has_centris_url = 0
    no_centris_url = 0
    has_address = 0
    no_address = 0
    
    print(f"\n🔍 Detailed Analysis:")
    
    for row in range(2, ws.max_row + 1):
        centris_cell = ws.cell(row=row, column=centris_col) if centris_col else None
        centris_url = centris_cell.hyperlink.target if centris_cell and centris_cell.hyperlink else None
        superficie = ws.cell(row=row, column=superficie_col).value if superficie_col else None
        address = ws.cell(row=row, column=address_col).value if address_col else None
        
        # Count addresses
        if address and address != "None":
            has_address += 1
        else:
            no_address += 1
        
        # Count Centris URLs
        if centris_url:
            has_centris_url += 1
        else:
            no_centris_url += 1
        
        # Count SuperficieDuterrain values
        if superficie and superficie != "":
            filled_superficie += 1
        else:
            empty_superficie += 1
    
    print(f"   Total listings: {total_rows}")
    print(f"   Listings with addresses: {has_address}")
    print(f"   Listings without addresses: {no_address}")
    print(f"   Listings with Centris URLs: {has_centris_url}")
    print(f"   Listings without Centris URLs: {no_centris_url}")
    print(f"   SuperficieDuterrain filled: {filled_superficie}")
    print(f"   SuperficieDuterrain empty: {empty_superficie}")
    
    # Show examples of filled SuperficieDuterrain
    print(f"\n✅ Examples of listings WITH SuperficieDuterrain:")
    count = 0
    for row in range(2, ws.max_row + 1):
        superficie = ws.cell(row=row, column=superficie_col).value if superficie_col else None
        address = ws.cell(row=row, column=address_col).value if address_col else None
        
        if superficie and superficie != "":
            print(f"   Row {row}: {superficie} | {address}")
            count += 1
            if count >= 5:
                break
    
    # Show examples of empty SuperficieDuterrain with URLs
    print(f"\n❌ Examples of listings WITHOUT SuperficieDuterrain (but have Centris URLs):")
    count = 0
    for row in range(2, ws.max_row + 1):
        centris_cell = ws.cell(row=row, column=centris_col) if centris_col else None
        centris_url = centris_cell.hyperlink.target if centris_cell and centris_cell.hyperlink else None
        superficie = ws.cell(row=row, column=superficie_col).value if superficie_col else None
        address = ws.cell(row=row, column=address_col).value if address_col else None
        
        if centris_url and (not superficie or superficie == ""):
            print(f"   Row {row}: EMPTY | {address} | URL: {centris_url}")
            count += 1
            if count >= 5:
                break
    
    print(f"\n💡 Root Cause Analysis:")
    print(f"   The SuperficieDuterrain processing requires:")
    print(f"   1. A valid Centris URL (hyperlink) - {has_centris_url} listings have this")
    print(f"   2. The extended-parse.py script to be run - only {filled_superficie} processed so far")
    print(f"   3. The script navigates to each Centris page to extract terrain size")
    
    print(f"\n🔧 Solutions:")
    if empty_superficie > 0 and has_centris_url > filled_superficie:
        missing_count = has_centris_url - filled_superficie
        print(f"   1. Run 'python3 extended-parse.py' and select 'SuperficieDuterrain'")
        print(f"   2. This will process {missing_count} listings with Centris URLs")
        print(f"   3. The process takes time as it visits each Centris property page")
        print(f"   4. Make sure your .env has correct USER_DATA_DIR configured")
    
    if no_address > 0:
        print(f"   5. {no_address} listings have no addresses - fix data extraction first")

if __name__ == "__main__":
    analyze_superficie_issue()
