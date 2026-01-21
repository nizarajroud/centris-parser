import boto3
import sys
import os
from dotenv import load_dotenv

load_dotenv()

if len(sys.argv) != 2:
    print("Usage: python update_bedrock_kb.py <new_source_url>")
    sys.exit(1)

new_url = sys.argv[1]

if not new_url or new_url.strip() == "":
    print("ℹ️  Empty URL provided, nothing to do")
    sys.exit(0)

kb_id = os.getenv('BEDROCK_KB_ID')
data_source_id = os.getenv('BEDROCK_DATA_SOURCE_ID')
region = os.getenv('AWS_DEFAULT_REGION', 'us-east-1')
profile = os.getenv('AWS_PROFILE')

session = boto3.Session(profile_name=profile, region_name=region)
client = session.client('bedrock-agent')

current = client.get_data_source(knowledgeBaseId=kb_id, dataSourceId=data_source_id)
existing_urls = current['dataSource']['dataSourceConfiguration']['webConfiguration']['sourceConfiguration']['urlConfiguration']['seedUrls']

if any(url['url'] == new_url for url in existing_urls):
    print(f"ℹ️  URL already exists in data source: {new_url}")
    sys.exit(0)

existing_urls.append({'url': new_url})

response = client.update_data_source(
    knowledgeBaseId=kb_id,
    dataSourceId=data_source_id,
    name='knowledge-base-quick-start-u6y2h-data-source',
    dataSourceConfiguration={
        'type': 'WEB',
        'webConfiguration': {
            'crawlerConfiguration': {
                'crawlerLimits': {
                    'rateLimit': 300
                }
            },
            'sourceConfiguration': {
                'urlConfiguration': {
                    'seedUrls': existing_urls
                }
            }
        }
    }
)

print(f"✅ Added URL to data source {data_source_id}: {new_url}")

# Start ingestion job to sync the data source
print("🔄 Starting data source synchronization...")
ingestion_response = client.start_ingestion_job(
    knowledgeBaseId=kb_id,
    dataSourceId=data_source_id
)

job_id = ingestion_response['ingestionJob']['ingestionJobId']
print(f"✅ Synchronization started with job ID: {job_id}")

# Update Excel file's details-page column
print("📝 Updating Excel file...")
import sqlite3
from openpyxl import load_workbook

extraction_file = os.getenv('EXTRACTION_CENTRIS', 'extraction-centris.xlsx')
db_path = os.getenv('CENTRIS_DB_PATH', 'centris.db')

# Get Centris ID from database for this URL
conn = sqlite3.connect(db_path)
cursor = conn.cursor()
cursor.execute('SELECT "Centris ID" FROM manual_details WHERE "Details page" = ?', (new_url,))
result = cursor.fetchone()
conn.close()

if result:
    centris_id = result[0]
    wb = load_workbook(extraction_file)
    ws = wb.active
    
    # Find the row with this Centris ID and update details-page column
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value) == str(centris_id):
            cell = ws.cell(row=row, column=2)
            cell.hyperlink = new_url
            cell.value = "More details"
            cell.style = "Hyperlink"
            wb.save(extraction_file)
            print(f"✅ Updated Excel file for Centris ID {centris_id}")
            break
else:
    print("⚠️  Could not find Centris ID for this URL in database")
