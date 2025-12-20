import os
from openpyxl import load_workbook
from dotenv import load_dotenv

load_dotenv()

extraction_file = os.getenv('EXTRACTION_CENTRIS', 'extraction-centris.xlsx')
print(f"Testing Excel file: {extraction_file}")

try:
    wb = load_workbook(extraction_file)
    ws = wb.active
    
    # Find WalkScore column
    walkscore_col = None
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header == "WalkScore":
            walkscore_col = col
            break
    
    print(f"WalkScore column: {walkscore_col}")
    
    # Check current values
    for row in range(2, min(10, ws.max_row + 1)):
        current_value = ws.cell(row=row, column=walkscore_col).value if walkscore_col else None
        address = ws.cell(row=row, column=2).value
        print(f"Row {row}: Address='{address}', WalkScore='{current_value}'")
    
    # Test writing a value
    if walkscore_col:
        ws.cell(row=2, column=walkscore_col, value=99)
        wb.save(extraction_file)
        print("Test write successful")
        
        # Read it back
        wb2 = load_workbook(extraction_file)
        ws2 = wb2.active
        test_value = ws2.cell(row=2, column=walkscore_col).value
        print(f"Test value read back: {test_value}")
        
except Exception as e:
    print(f"Error: {e}")
